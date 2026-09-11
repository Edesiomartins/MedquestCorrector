from io import BytesIO

import pymupdf as fitz
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.services.discursive_import.canonical_pdf import add_student_identity_header


def _two_page_pdf() -> bytes:
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, _h = A4
    for page_number, question_number in enumerate((38, 39)):
        c.drawString(42, 700, f"QUESTÃO {question_number}")
        y = 650
        for _ in range(4):
            c.line(42, y, w - 42, y)
            y -= 24
        if page_number == 0:
            c.showPage()
    c.save()
    return buf.getvalue()


def _long_horizontal_line_count(page: fitz.Page) -> int:
    width = float(page.rect.width)
    count = 0
    for drawing in page.get_drawings():
        for item in drawing.get("items") or []:
            if not item or item[0] != "l":
                continue
            p1, p2 = item[1], item[2]
            if abs(float(p1.y) - float(p2.y)) <= 2 and abs(float(p2.x) - float(p1.x)) >= width * 0.55:
                count += 1
    return count


def test_identity_header_is_added_to_every_page_without_changing_page_size():
    raw = _two_page_pdf()
    stamped = add_student_identity_header(raw)
    doc = fitz.open(stream=stamped, filetype="pdf")
    try:
        assert doc.page_count == 2
        for page in doc:
            text = page.get_text()
            assert "Nome:" in text
            assert "Matrícula:" in text
            assert round(page.rect.width) == round(A4[0])
            assert round(page.rect.height) == round(A4[1])
    finally:
        doc.close()


def test_canonical_pdf_has_at_least_six_writing_lines_when_space_exists():
    from docx import Document

    from app.services.discursive_import.layout_detector import normalize_docx_to_pdf

    doc = Document()
    doc.add_paragraph("Questão 38")
    doc.add_paragraph("Enunciado curto o bastante para sobrar espaço vertical.")
    for _ in range(4):
        doc.add_paragraph("_" * 90)
    buf = BytesIO()
    doc.save(buf)

    normalized = normalize_docx_to_pdf(buf.getvalue())
    stamped = add_student_identity_header(normalized["canonical_pdf"])
    pdf = fitz.open(stream=stamped, filetype="pdf")
    try:
        assert all(_long_horizontal_line_count(page) >= 6 for page in pdf)
    finally:
        pdf.close()
