from io import BytesIO

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.services.discursive_import.layout_detector import (
    build_template_manifest,
    detect_pdf_layout,
    normalize_docx_to_pdf,
    validate_confirmed_layout,
)


def _pdf_with_lined_question(number=38, second_number=None):
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setFont('Helvetica-Bold', 12)
    c.drawString(42, h - 70, f'Questão {number}')
    c.setFont('Helvetica', 10)
    c.drawString(42, h - 92, 'Explique a diferença morfológica fundamental entre as articulações.')
    y = h - 145
    for _ in range(6):
        c.line(42, y, w - 42, y)
        y -= 24
    if second_number is not None:
        c.setFont('Helvetica-Bold', 12)
        c.drawString(42, y - 12, f'Questão {second_number}')
        c.setFont('Helvetica', 10)
        c.drawString(42, y - 34, 'Defina o conceito solicitado.')
        y2 = y - 70
        for _ in range(4):
            c.line(42, y2, w - 42, y2)
            y2 -= 24
    c.save()
    return buf.getvalue()


def test_detects_question_38_and_lined_answer_area():
    result = detect_pdf_layout(_pdf_with_lined_question())
    assert [q['question_number'] for q in result['questions']] == [38]
    q = result['questions'][0]
    assert q['page_index'] == 0
    assert q['height_pt'] > 80
    assert q['width_pt'] > 400
    assert q['provenance'] == 'answer_lines'
    assert q['confidence'] >= 0.9
    assert 'diferença morfológica' in q['question_text'].lower()
    assert result['pages'][0]['preview_data_url'].startswith('data:image/png;base64,')


def test_detects_non_contiguous_question_numbers():
    result = detect_pdf_layout(_pdf_with_lined_question(38, 41))
    assert [q['question_number'] for q in result['questions']] == [38, 41]


def test_build_template_manifest_repeats_template_pages():
    detected = detect_pdf_layout(_pdf_with_lined_question())
    manifest = build_template_manifest('exam-123', detected['pages'], detected['questions'])
    assert manifest['template_repeat'] is True
    assert manifest['template_page_count'] == 1
    assert manifest['source'] == 'external_discursive'
    assert manifest['pages'][0]['boxes'][0]['question_number'] == 38
    assert manifest['pages'][0]['student_id'] == ''


def test_validation_rejects_duplicate_question_numbers():
    pages = [{'page_index': 0, 'width_pt': 595.0, 'height_pt': 842.0}]
    questions = [
        {'question_number': 38, 'page_index': 0, 'question_text': 'A', 'x_pt': 20, 'y_bottom_pt': 20, 'width_pt': 200, 'height_pt': 100},
        {'question_number': 38, 'page_index': 0, 'question_text': 'B', 'x_pt': 20, 'y_bottom_pt': 140, 'width_pt': 200, 'height_pt': 100},
    ]
    errors, _warnings = validate_confirmed_layout(pages, questions)
    assert any('duplicado' in item.lower() for item in errors)


def test_validation_rejects_box_outside_page():
    pages = [{'page_index': 0, 'width_pt': 595.0, 'height_pt': 842.0}]
    questions = [
        {'question_number': 1, 'page_index': 0, 'question_text': 'A', 'x_pt': 500, 'y_bottom_pt': 20, 'width_pt': 200, 'height_pt': 100},
    ]
    errors, _warnings = validate_confirmed_layout(pages, questions)
    assert any('limites da página' in item.lower() for item in errors)


def test_docx_is_normalized_to_canonical_pdf_with_real_question_number():
    from docx import Document

    doc = Document()
    doc.add_paragraph('Curso de Medicina')
    doc.add_paragraph('Questão 38')
    doc.add_paragraph('O punho é composto por oito ossos. Analise a diferença morfológica.')
    for _ in range(6):
        doc.add_paragraph('_' * 90)
    buf = BytesIO()
    doc.save(buf)

    normalized = normalize_docx_to_pdf(buf.getvalue())
    assert normalized['canonical_pdf'][:4] == b'%PDF'
    assert normalized['warnings']
    detected = detect_pdf_layout(normalized['canonical_pdf'])
    assert [q['question_number'] for q in detected['questions']] == [38]
    assert detected['questions'][0]['provenance'] == 'answer_lines'


def test_blank_space_without_lines_is_proposed_for_manual_confirmation():
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    c.setFont('Helvetica-Bold', 12)
    c.drawString(42, h - 70, 'Q 7')
    c.setFont('Helvetica', 10)
    c.drawString(42, h - 92, 'Descreva o mecanismo fisiológico solicitado.')
    c.save()
    result = detect_pdf_layout(buf.getvalue())
    assert result['questions'][0]['question_number'] == 7
    assert result['questions'][0]['provenance'] == 'blank_space'
    assert result['questions'][0]['height_pt'] > 100


def test_multiple_pages_preserve_page_indexes():
    buf = BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    w, h = A4
    for number in (38, 39):
        c.setFont('Helvetica-Bold', 12)
        c.drawString(42, h - 70, f'QUESTÃO {number}')
        c.setFont('Helvetica', 10)
        c.drawString(42, h - 92, f'Enunciado da questão {number}.')
        y = h - 145
        for _ in range(5):
            c.line(42, y, w - 42, y)
            y -= 24
        if number == 38:
            c.showPage()
    c.save()
    result = detect_pdf_layout(buf.getvalue())
    assert [(q['question_number'], q['page_index']) for q in result['questions']] == [(38, 0), (39, 1)]
    assert len(result['pages']) == 2


def test_overlap_is_warning_not_error():
    pages = [{'page_index': 0, 'width_pt': 595.0, 'height_pt': 842.0}]
    questions = [
        {'question_number': 1, 'page_index': 0, 'question_text': 'A', 'x_pt': 20, 'y_bottom_pt': 20, 'width_pt': 200, 'height_pt': 100},
        {'question_number': 2, 'page_index': 0, 'question_text': 'B', 'x_pt': 30, 'y_bottom_pt': 30, 'width_pt': 200, 'height_pt': 100},
    ]
    errors, warnings = validate_confirmed_layout(pages, questions)
    assert errors == []
    assert any('sobrepõem' in item for item in warnings)


def test_docx_question_label_with_underscores_does_not_pollute_question_text():
    from docx import Document
    doc = Document()
    doc.add_paragraph('Questão 38 ' + '_' * 50)
    doc.add_paragraph('O punho é composto por oito ossos carpianos.')
    doc.add_paragraph('Analise a diferença morfológica fundamental.')
    for _ in range(6):
        doc.add_paragraph('_' * 90)
    buf = BytesIO(); doc.save(buf)
    normalized = normalize_docx_to_pdf(buf.getvalue())
    detected = detect_pdf_layout(normalized['canonical_pdf'])
    q = detected['questions'][0]
    assert 'punho' in q['question_text'].lower()
    assert not q['question_text'].lstrip().startswith('_')
